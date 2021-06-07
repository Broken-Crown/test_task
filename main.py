import openpyxl
import random
from pathlib import Path
from datetime import datetime, timedelta

_COMPANY_NAME: list[str] = ["Компания_1", "Компания_2", "Компания_3"]
_OPERATIONS: list[str] = ["Выплата зп", "Отплата налогов", "Закупка оборудования"]
_TABLE_HEADER: list[str] = ["Компания", "Дата", "Тип операции", "Сумма"]
_CUSTOM_FIRST_TABLE_PATH: Path = Path(r".\first_table.xlsx")
_CUSTOM_SECOND_TABLE_PATH: Path = Path(r".\second_table.xlsx")
_MAX_ROW_NUMBER: int = 10000
_MAX_COL_NUMBER: int = len(_TABLE_HEADER)


def create_table():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Информация по операциям"
    for index, cell in enumerate(ws.iter_cols(max_row=1, max_col=_MAX_COL_NUMBER)):
        cell[0].value = _TABLE_HEADER[index]

    current_date = datetime.now()
    for row in ws.iter_rows(min_row=2, max_row=_MAX_ROW_NUMBER, max_col=_MAX_COL_NUMBER):
        row[0].value = random.choice(_COMPANY_NAME)
        row[1].value = f"{current_date.year}-{current_date.month:02}-{current_date.day:02} " \
                       f"{current_date.hour:02}:{current_date.month:02}:{current_date.second:02}"
        row[2].value = random.choice(_OPERATIONS)
        row[3].value = random.randint(100, 500000)

        current_date += timedelta(0, random.randint(0, 10))

    wb.save(_CUSTOM_FIRST_TABLE_PATH)


if __name__ == "__main__":
    print("Опять работать?")
    create_table()
    print("Дело сделано!")
